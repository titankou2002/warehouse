#!/usr/bin/env python3
"""Analyze 永安倉庫 Excel: SKU-row vs batch-row qty patterns for box/piece migration.
Run on BIGT-MACMINI (or after CopyToBox of the xlsx).
"""
from __future__ import annotations
import json, re, sys
from collections import Counter, defaultdict
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles.colors import COLOR_INDEX
except ImportError:
    print("Need openpyxl: pip install openpyxl"); sys.exit(1)

DEFAULT = Path("/Users/titankou2002/Library/CloudStorage/GoogleDrive-titankou2002@gmail.com/我的雲端硬碟/BT/Antigravity/倉庫全集/永安倉庫2025庫位表 (2).xlsx")

SKU_RE = re.compile(r"^[A-Z]{1,3}\d{3,}[A-Z0-9\-]*$", re.I)
BATCH_RE = re.compile(r"^[A-Z]{1,3}\d{1,4}[A-Z]?$", re.I)  # e.g. GD6, GB4
NUM_RE = re.compile(r"^\s*-?\d+(\.\d+)?\s*$")

def cell_rgb(cell):
    fill = cell.fill
    if not fill or not fill.fgColor: return None
    c = fill.fgColor
    if c.type == "rgb" and c.rgb: return str(c.rgb)[-6:].upper()
    if c.type == "theme": return f"THEME:{c.theme}"
    if c.type == "indexed": return f"IDX:{c.indexed}"
    return None

def font_rgb(cell):
    f = cell.font
    if not f or not f.color: return None
    c = f.color
    if c.type == "rgb" and c.rgb: return str(c.rgb)[-6:].upper()
    if c.type == "theme": return f"THEME:{c.theme}"
    return None

def is_num(v):
    if v is None or v == "": return False
    if isinstance(v, (int, float)) and not isinstance(v, bool): return True
    return bool(NUM_RE.match(str(v).strip()))

def as_num(v):
    if v is None or v == "": return None
    if isinstance(v, (int, float)) and not isinstance(v, bool): return float(v)
    s = str(v).strip()
    try: return float(s)
    except: return None

def classify_text(v):
    if v is None: return "empty"
    s = str(v).strip()
    if not s: return "empty"
    if is_num(v): return "number"
    if SKU_RE.match(s): return "sku_like"
    if BATCH_RE.match(s): return "batch_like"
    if any(x in s for x in ("箱", "片", "排", "列", "層", "區")): return "label"
    return "text"

def analyze_sheet(ws, max_rows=400, max_cols=80):
    """Heuristic scan: look for adjacent row pairs that look like SKU + batch."""
    patterns = Counter()
    examples = []
    color_counts = Counter()
    # Sample raw grid for docs
    sample_grid = []
    for r in range(1, min(25, ws.max_row or 1) + 1):
        row = []
        for c in range(1, min(20, ws.max_column or 1) + 1):
            cell = ws.cell(r, c)
            row.append({"v": cell.value, "bg": cell_rgb(cell), "fg": font_rgb(cell)})
        sample_grid.append(row)

    # Pairwise scan
    mr = min(ws.max_row or 1, max_rows)
    mc = min(ws.max_column or 1, max_cols)
    for r in range(1, mr):
        for c in range(1, mc + 1):
            a = ws.cell(r, c)
            b = ws.cell(r + 1, c)
            ta, tb = classify_text(a.value), classify_text(b.value)
            # Candidate: SKU-like over batch-like, or SKU over number/text
            if ta == "sku_like" and tb in ("batch_like", "text", "number", "empty"):
                # Look for qty in same column or nearby columns on both rows
                sku_qty = None; batch_qty = None
                sku_qty_cell = None; batch_qty_cell = None
                for dc in range(0, 4):
                    if c + dc > mc: break
                    av = ws.cell(r, c + dc).value
                    bv = ws.cell(r + 1, c + dc).value
                    if sku_qty is None and is_num(av) and dc > 0:
                        sku_qty = as_num(av); sku_qty_cell = (r, c + dc)
                    if batch_qty is None and is_num(bv):
                        # number on batch row (same or adjacent col)
                        if dc >= 0:
                            batch_qty = as_num(bv); batch_qty_cell = (r + 1, c + dc)
                # Also check cells immediately right of each
                if sku_qty is None and c + 1 <= mc and is_num(ws.cell(r, c + 1).value):
                    sku_qty = as_num(ws.cell(r, c + 1).value); sku_qty_cell = (r, c + 1)
                if batch_qty is None and c + 1 <= mc and is_num(ws.cell(r + 1, c + 1).value):
                    batch_qty = as_num(ws.cell(r + 1, c + 1).value); batch_qty_cell = (r + 1, c + 1)

                has_sku = sku_qty is not None
                has_batch = batch_qty is not None
                if has_sku and not has_batch: pat = "a_qty_only_sku_row"
                elif has_batch and not has_sku: pat = "b_qty_only_batch_row"
                elif has_sku and has_batch: pat = "c_both_rows"
                else: pat = "d_ambiguous_or_no_qty"
                patterns[pat] += 1
                if len(examples) < 40 or (pat != "d_ambiguous_or_no_qty" and sum(1 for e in examples if e["pattern"]==pat) < 8):
                    examples.append({
                        "sheet": ws.title,
                        "sku_cell": f"{a.coordinate}",
                        "sku": str(a.value),
                        "batch_cell": f"{b.coordinate}",
                        "batch": None if b.value is None else str(b.value),
                        "sku_qty": sku_qty,
                        "sku_qty_cell": None if not sku_qty_cell else f"R{sku_qty_cell[0]}C{sku_qty_cell[1]}",
                        "batch_qty": batch_qty,
                        "batch_qty_cell": None if not batch_qty_cell else f"R{batch_qty_cell[0]}C{batch_qty_cell[1]}",
                        "sku_bg": cell_rgb(a), "batch_bg": cell_rgb(b),
                        "sku_fg": font_rgb(a), "batch_fg": font_rgb(b),
                        "pattern": pat,
                        # Current GAS heuristic (from task): blank SKU qty + number on batch → BoxQty
                        "gas_would_set_boxqty_from_batch": (not has_sku) and has_batch,
                    })
            bg = cell_rgb(a)
            if bg: color_counts[bg] += 1
    return {
        "title": ws.title,
        "dims": [ws.max_row, ws.max_column],
        "patterns": dict(patterns),
        "color_counts_top": color_counts.most_common(15),
        "sample_grid_25x20": sample_grid,
        "examples": examples,
    }

def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("excel_qty_analysis.json")
    print("Opening", path)
    wb = openpyxl.load_workbook(path, data_only=True)
    report = {"file": str(path), "sheets": wb.sheetnames, "analyses": []}
    print("Sheets:", wb.sheetnames)
    for name in wb.sheetnames:
        # skip obvious non-zone sheets
        if any(x in name for x in ("說明", "README", "價目", "目錄", "DataTable", "LOG", "log", "History")):
            print(" skip", name); continue
        ws = wb[name]
        print(" analyzing", name, ws.max_row, "x", ws.max_column)
        report["analyses"].append(analyze_sheet(ws))
    # Aggregate
    agg = Counter()
    all_ex = []
    for a in report["analyses"]:
        agg.update(a["patterns"])
        all_ex.extend(a["examples"])
    report["aggregate_patterns"] = dict(agg)
    # Prefer 20 diverse examples
    picked = []
    for want in ("a_qty_only_sku_row", "b_qty_only_batch_row", "c_both_rows", "d_ambiguous_or_no_qty"):
        picked.extend([e for e in all_ex if e["pattern"] == want][:5])
    report["sample_20"] = picked[:20]
    out.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print("Wrote", out)
    print("AGG", dict(agg))
    print("SAMPLE count", len(report["sample_20"]))

if __name__ == "__main__":
    main()
